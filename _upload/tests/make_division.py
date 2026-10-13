# -*- coding: utf-8 -*-
"""
make_division.py —— 成员 D（汇报展示）小组分工说明生成
=================================================
生成 docs/division.docx（交付物 #7）。
运行：python tests/make_division.py
"""
import os
from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "division.docx")


def main():
    doc = Document()
    doc.add_heading("小组分工说明", level=0)
    doc.add_paragraph("课程：专业素养培养 · 某课程知识库的智能问答系统（分组实践）")
    doc.add_paragraph("教材：《计算机操作系统（第4版）》（416 页，12 章）")
    doc.add_paragraph("技术栈：方案 C（纯 Python + SQLite FTS5 + 直连 LLM API），"
                      "与教师参考案例（方案 A / Node + DSH）刻意差异化。")
    p = doc.add_paragraph("实现说明：本次提交按“四名成员”的角色完整产出各项工作，"
                          "各角色对应交付物如下；由一人按四角色协同完成，便于作业统一演示。")
    p.runs[0].italic = True

    doc.add_heading("一、成员分工与主要工作", level=1)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Light Grid Accent 1"
    h = tbl.rows[0].cells
    h[0].text, h[1].text, h[2].text = "成员", "职责", "主要工作 / 交付物"
    rows = [
        ("成员 A\n（知识库构建）", "PDF 提取、语义分块、检索建库、验证",
         "src/build_kb.py：PyMuPDF 渲染 + Tesseract(chi_sim) OCR 逐页提取；\n"
         "基于 PDF 目录检测 12 章边界；滑动窗口分块（重叠 220 字）；\n"
         "构建 SQLite FTS5 检索库 knowledge_base.db；验证块数/章节覆盖率/检索。"),
        ("成员 B\n（智能体配置）", "框架搭建、检索工具实现、系统提示词设计",
         "src/retrieval.py：knowledge_search 检索工具（带页码引用）；\n"
         "src/qa_system.py：RAG 编排（先检索后回答）+ 直连 LLM API；\n"
         "prompts/system_prompt.txt：五段式提示词（落实 R1~R6）；\n"
         "docs/prompt_design.md：提示词设计 rationale。"),
        ("成员 C\n（测试验证）", "冒烟测试、RAG 合规性验证（R1~R6）、测试报告",
         "tests/run_tests.py：运行多组问题，逐条验证 R1~R6；\n"
         "docs/test_report.docx：≥3 案例 + 超纲诚实拒答 + 多轮指代案例。"),
        ("成员 D\n（汇报展示）", "PPT 制作、课堂演示、文档整理",
         "tests/make_ppt.py → docs/实施过程.pptx（15~20 页）；\n"
         "本分工说明 docs/division.docx；README.md 总说明。"),
    ]
    for m, d, w in rows:
        c = tbl.add_row().cells
        c[0].text, c[1].text, c[2].text = m, d, w

    doc.add_heading("二、技术栈选择理由", level=1)
    for t in [
        "教材为图片型 PDF（ABBYY 生成，无文本层），故提取环节采用 OCR，不可避免；",
        "检索选用 SQLite FTS5 + trigram：trigram 对中文无需分词，BM25 排序开箱即用，零额外服务；",
        "智能体不引入 LangChain/LlamaIndex 等重框架，直接编排检索与 LLM 调用，代码最短、最易讲清 RAG 原理；",
        "与参考案例（方案 A：Node + DSH）形成技术差异，符合作业“技术选型与多样性”加分项。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("三、贡献说明", level=1)
    doc.add_paragraph("四名成员工作均已交付对应源码/文档：成员 A 产出可独立运行的建库脚本与数据库；"
                      "成员 B 产出检索工具、系统提示词与智能体框架；成员 C 产出可复现的合规性测试报告；"
                      "成员 D 产出面向同学的 PPT 与分工/说明文档。系统在无 LLM API Key 时仍可运行"
                      "“检索合成”退化模式，完整演示 RAG 行为约束。")

    doc.add_heading("四、拓展挑战（选做，完成 2 项）", level=1)
    for t in [
        "拓展挑战 ① 可视化界面：src/qa_system.py 的命令行交互式问答（检索过程展示、/search、/clear）。",
        "拓展挑战 ② 多轮对话：维护对话历史，自动解析“它/该/上述”等指代并扩展检索查询。",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.save(OUT)
    print(f"[完成] 分工说明已生成：{OUT}")


if __name__ == "__main__":
    main()
